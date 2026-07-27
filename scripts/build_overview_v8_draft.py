#!/usr/bin/env python3
"""Build a motivation-first v8 overview draft from the original PPT layout."""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from build_overview_compliant import (
    GOLD,
    GOLD_PALE,
    GRID,
    INK,
    MIN_PT,
    MUTED,
    RED,
    RED_PALE,
    SLATE,
    TEAL,
    TEAL_PALE,
    WHITE,
    add_arrow,
    add_box,
    add_dot,
    add_formula,
    add_matrix,
    add_rule,
    add_text,
    set_cell_fill,
    set_line,
)


W = 6.86
H = 2.60


def behavior_icon(slide, x, y, color, fork=False):
    if fork:
        add_dot(slide, x + 0.08, y, 0.09, color, color)
        add_dot(slide, x, y + 0.13, 0.09, color, color)
        add_dot(slide, x + 0.16, y + 0.13, 0.09, color, color)
        add_rule(slide, x + 0.12, y + 0.07, x + 0.04, y + 0.15, color, 0.7)
        add_rule(slide, x + 0.12, y + 0.07, x + 0.20, y + 0.15, color, 0.7)
    else:
        for i in range(3):
            add_dot(slide, x + i * 0.12, y + 0.06, 0.09, color, color)
        add_rule(slide, x + 0.09, y + 0.10, x + 0.15, y + 0.10, color, 0.7)
        add_rule(slide, x + 0.21, y + 0.10, x + 0.27, y + 0.10, color, 0.7)


def lock_icon(slide, x, y, color):
    body = add_box(slide, x, y + 0.07, 0.13, 0.11, color, color, False)
    body.line.fill.background()
    arc = slide.shapes.add_shape(MSO_SHAPE.ARC, Inches(x + 0.02), Inches(y), Inches(0.09), Inches(0.11))
    set_line(arc, color, 0.8)
    arc.fill.background()


def task_icons(slide, x, y):
    add_matrix(slide, x, y, 2, 2, GOLD, RGBColor(247, 232, 184))
    behavior_icon(slide, x + 0.27, y, TEAL, fork=True)
    lock_icon(slide, x + 0.65, y, RED)


def weight_marker(slide, x, y, color, pale):
    add_box(slide, x, y, 0.08, 0.19, color, color, False)
    add_box(slide, x + 0.10, y + 0.04, 0.08, 0.15, pale, pale, False)
    add_box(slide, x + 0.20, y + 0.09, 0.08, 0.10, pale, pale, False)


def build(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    # Keep the original horizontal reading order, but give the innovation
    # blocks the visual weight and reserve short annotations for the rest.
    add_text(slide, 0.08, 0.06, 1.30, 0.20, "WHY TWO SIGNALS?", size=9, color=RED, bold=True)
    add_text(slide, 1.42, 0.07, 0.82, 0.18, "01  INPUT", size=9, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 2.58, 0.07, 1.34, 0.18, "INNOVATION 1", size=9, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 4.08, 0.07, 1.52, 0.18, "INNOVATION 2", size=9, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 5.82, 0.07, 0.85, 0.18, "OUTPUT", size=9, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_rule(slide, 0.08, 0.30, W - 0.08, 0.30, SLATE, 1.1)

    # Motivation panel: two graphical counterexamples.
    add_box(slide, 0.08, 0.39, 1.22, 2.05, RED_PALE, RGBColor(241, 173, 163))
    add_text(slide, 0.16, 0.47, 1.06, 0.15, "confidence only", size=9, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_box(slide, 0.16, 0.68, 0.45, 0.25, WHITE, RGBColor(241, 173, 163))
    add_box(slide, 0.75, 0.68, 0.45, 0.25, WHITE, RGBColor(241, 173, 163))
    add_text(slide, 0.19, 0.72, 0.10, 0.11, "A", size=9, color=RED, bold=True)
    add_text(slide, 0.78, 0.72, 0.10, 0.11, "B", size=9, color=RED, bold=True)
    for x in (0.30, 0.89):
        for i in range(3):
            add_box(slide, x + i * 0.08, 0.77, 0.055, 0.055, RED if i == 1 else RGBColor(249, 206, 198), RED, False)
    behavior_icon(slide, 0.28, 1.00, TEAL, fork=False)
    behavior_icon(slide, 0.88, 1.00, RED, fork=True)
    add_text(slide, 0.15, 1.25, 1.08, 0.15, "same q ≠ behavior", size=9, color=TEAL, bold=True, align=PP_ALIGN.CENTER)

    add_text(slide, 0.16, 1.49, 1.06, 0.15, "majority only", size=9, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    for i in range(3):
        add_dot(slide, 0.28 + i * 0.13, 1.72, 0.09, TEAL, TEAL)
    add_dot(slide, 0.99, 1.72, 0.12, GOLD, GOLD)
    add_rule(slide, 0.68, 1.77, 0.88, 1.77, GRID, 0.9)
    add_text(slide, 0.15, 1.94, 1.08, 0.15, "count ≠ target risk", size=9, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 0.15, 2.24, 1.08, 0.16, "quality + structure", size=9, color=SLATE, bold=True, align=PP_ALIGN.CENTER)

    # Stage 1: show only the fixed target and the peer set.
    add_box(slide, 1.46, 0.47, 0.91, 0.37, GOLD_PALE, RGBColor(225, 193, 102))
    add_text(slide, 1.51, 0.56, 0.81, 0.14, "0  target", size=9, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    for i in range(4):
        add_box(slide, 1.54 + i * 0.12, 0.72, 0.07, 0.07, GOLD if i == 0 else RGBColor(247, 215, 137), GOLD, False)
    add_box(slide, 1.46, 1.05, 0.91, 0.37, TEAL_PALE, RGBColor(153, 206, 201))
    add_text(slide, 1.50, 1.14, 0.83, 0.14, "1…N−1 peers", size=9, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    for i in range(4):
        add_box(slide, 1.54 + i * 0.12, 1.31, 0.07, 0.07, TEAL, TEAL, False)
    add_text(slide, 1.48, 1.73, 0.88, 0.14, "N paths", size=9, color=MUTED, align=PP_ALIGN.CENTER)
    add_arrow(slide, 2.40, 1.17, 0.12, SLATE)

    # Innovation 1: parallel quality and task-geometry streams.
    add_box(slide, 2.53, 0.39, 1.40, 2.05, RED_PALE, RGBColor(241, 173, 163))
    add_text(slide, 2.61, 0.45, 1.24, 0.16, "quality × geometry", size=9, color=RED, bold=True, align=PP_ALIGN.CENTER)

    for i in range(5):
        add_box(slide, 2.64 + i * 0.08, 0.76, 0.06, 0.06, RED if i == 3 else RGBColor(249, 206, 198), RED, False)
    add_text(slide, 2.58, 0.88, 0.55, 0.13, "weakest-k", size=9, color=MUTED, align=PP_ALIGN.CENTER)
    add_rule(slide, 3.04, 0.79, 3.18, 0.79, RED, 0.8)
    add_dot(slide, 3.17, 0.72, 0.13, TEAL, TEAL)
    add_formula(slide, 2.96, 0.98, 0.80, 0.18, [[("q", None), ("i", "sub"), ("=c", None), ("i", "sub"), ("v", None), ("i", "sub")]], size=9)

    task_icons(slide, 2.54, 1.32)
    add_text(slide, 2.58, 1.57, 0.70, 0.13, "adapters", size=9, color=MUTED, align=PP_ALIGN.CENTER)
    add_matrix(slide, 3.36, 1.31, 2, 2, TEAL, RGBColor(190, 228, 224))
    add_formula(slide, 2.99, 1.70, 0.66, 0.18, [[("K=ΦΦ", None), ("T", "sup")]], size=9)
    add_rule(slide, 3.78, 1.00, 3.78, 1.27, RED, 0.8)
    add_rule(slide, 3.78, 1.27, 3.78, 1.55, TEAL, 0.8)
    add_text(slide, 3.69, 1.93, 0.18, 0.13, "w", size=9, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    weight_marker(slide, 3.64, 1.68, GOLD, RGBColor(246, 232, 183))
    add_arrow(slide, 3.96, 1.16, 0.12, SLATE)

    # Innovation 2: target/peer mass, paired spectra, and normalized contrast.
    add_box(slide, 4.06, 0.39, 1.66, 2.05, TEAL_PALE, RGBColor(153, 206, 201))
    add_text(slide, 4.14, 0.45, 1.50, 0.16, "target-conditioned contrast", size=9, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_box(slide, 4.15, 0.70, 0.47, 0.25, GOLD_PALE, RGBColor(225, 193, 102))
    add_text(slide, 4.18, 0.76, 0.41, 0.13, "0  α", size=9, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    for i in range(4):
        add_dot(slide, 4.80 + i * 0.13, 0.78, 0.09, TEAL, TEAL)
    add_text(slide, 4.78, 0.94, 0.60, 0.13, "peers  β", size=9, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_formula(slide, 4.18, 1.07, 1.37, 0.30, [[("w", None), ("i", "sub"), ("=q", None), ("i", "sub"), ("/Q", None)], [("π", None), ("j", "sub"), ("=q", None), ("j", "sub"), ("/Q", None), ("−0", "sub")]], size=9)

    add_text(slide, 4.14, 1.45, 0.26, 0.13, "mix", size=9, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    add_matrix(slide, 4.43, 1.40, 2, 3, GOLD, RGBColor(246, 232, 183))
    weight_marker(slide, 5.06, 1.40, GOLD, RGBColor(246, 232, 183))
    add_text(slide, 4.14, 1.75, 0.26, 0.13, "peer", size=9, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_matrix(slide, 4.43, 1.70, 2, 3, TEAL, RGBColor(190, 228, 224))
    weight_marker(slide, 5.06, 1.70, TEAL, RGBColor(190, 228, 224))
    add_formula(slide, 4.18, 2.06, 1.35, 0.20, [[("χ", None), ("0", "sub"), ("=S", None), ("mix", "sub"), ("−βS", None), ("peer", "sub")]], size=9)
    add_arrow(slide, 5.70, 1.18, 0.12, SLATE)

    # Output column: risk and independent selection, plus the small boundary note.
    add_box(slide, 5.83, 0.39, 0.94, 1.01, GOLD_PALE, RGBColor(225, 193, 102))
    add_text(slide, 5.88, 0.46, 0.84, 0.14, "fixed-target risk", size=9, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    add_formula(slide, 5.88, 0.65, 0.84, 0.30, [[("R", None), ("CI", "sub"), ("=clip", None), ("[0,1]", "sub")], [("χ", None), ("0", "sub"), ("/H", None), ("2", "sub"), ("(α)", None)]], size=9)
    add_rule(slide, 6.00, 1.04, 6.55, 1.04, GOLD, 0.9)
    add_dot(slide, 5.96, 1.00, 0.08, GOLD, GOLD)
    add_dot(slide, 6.51, 1.00, 0.08, TEAL, TEAL)
    add_box(slide, 5.83, 1.40, 0.94, 0.86, TEAL_PALE, RGBColor(153, 206, 201))
    add_text(slide, 5.88, 1.46, 0.84, 0.14, "self-consistent", size=9, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_formula(slide, 5.87, 1.68, 0.86, 0.18, [[("i", None), ("*", "sup"), ("=argmax", None), ("i", "sub")]], size=9)
    add_formula(slide, 5.87, 1.86, 0.86, 0.18, [[("q", None), ("i", "sub"), ("(Kw)", None), ("i", "sub")]], size=9)
    add_box(slide, 5.91, 1.12, 0.78, 0.24, WHITE, RGBColor(225, 193, 102))
    add_formula(
        slide,
        5.95,
        1.12,
        0.70,
        0.24,
        [
            [("N=1: 1−q", None), ("0", "sub")],
            [("q", None), ("0", "sub"), ("Q", None), ("−0", "sub"), ("=0: 1", None)],
        ],
        size=9,
    )

    # Production notes stay on a separate slide and are not part of the figure export.
    notes = prs.slides.add_slide(blank)
    notes.background.fill.solid()
    notes.background.fill.fore_color.rgb = WHITE
    add_text(notes, 0.35, 0.35, 6.1, 0.35, "AAAI-27 production notes", size=14, bold=True)
    add_text(
        notes,
        0.40,
        0.95,
        6.0,
        1.1,
        "Final canvas: 17.42 × 6.60 cm\nFigure labels and formulas: 9 pt minimum at final size\nKeep slide 1 at 100% when exporting; use vector PDF or ≥300 dpi raster",
        size=10,
        color=MUTED,
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.output)


if __name__ == "__main__":
    main()
