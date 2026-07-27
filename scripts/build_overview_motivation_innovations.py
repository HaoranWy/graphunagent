#!/usr/bin/env python3
"""Build the motivation-first, innovation-focused PWSC-CI overview slide."""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from build_overview_compliant import (
    GOLD,
    GOLD_PALE,
    GRID,
    INK,
    MIN_PT,
    MUTED,
    RED,
    RED_PALE,
    SLATE,
    TEAL,
    TEAL_PALE,
    WHITE,
    add_arrow,
    add_box,
    add_dot,
    add_formula,
    add_matrix,
    add_rule,
    add_text,
    set_cell_fill,
    set_line,
)


W = 6.86
H = 2.60


def add_icon_node(slide, x, y, color, *, fork=False):
    """Draw a compact behavior graph used in the motivation panel."""
    if fork:
        add_dot(slide, x + 0.09, y, 0.10, color, color)
        add_dot(slide, x, y + 0.14, 0.10, color, color)
        add_dot(slide, x + 0.18, y + 0.14, 0.10, color, color)
        add_rule(slide, x + 0.14, y + 0.08, x + 0.05, y + 0.17, color, 0.8)
        add_rule(slide, x + 0.14, y + 0.08, x + 0.23, y + 0.17, color, 0.8)
    else:
        for i in range(3):
            add_dot(slide, x + i * 0.15, y + 0.07, 0.10, color, color)
        add_rule(slide, x + 0.10, y + 0.12, x + 0.15, y + 0.12, color, 0.8)
        add_rule(slide, x + 0.25, y + 0.12, x + 0.30, y + 0.12, color, 0.8)


def add_lock_icon(slide, x, y, color):
    body = add_box(slide, x, y + 0.07, 0.14, 0.12, color, color, False)
    body.line.fill.background()
    arc = slide.shapes.add_shape(MSO_SHAPE.ARC, Inches(x + 0.025), Inches(y), Inches(0.09), Inches(0.12))
    set_line(arc, color, 0.9)
    arc.fill.background()


def add_domain_icons(slide, x, y):
    # Three small, visual task adapters; labels are intentionally omitted.
    add_matrix(slide, x, y + 0.01, 2, 2, GOLD, RGBColor(247, 232, 184))
    add_icon_node(slide, x + 0.30, y + 0.01, TEAL, fork=True)
    add_lock_icon(slide, x + 0.68, y + 0.01, RED)


def add_weight_bar(slide, x, y, color, height, *, pale):
    add_box(slide, x, y + (0.22 - height), 0.10, height, color, color, False)
    add_box(slide, x + 0.12, y + (0.22 - 0.16), 0.10, 0.16, pale, pale, False)
    add_box(slide, x + 0.24, y + (0.22 - 0.10), 0.10, 0.10, pale, pale, False)


def add_risk_gauge(slide, x, y):
    add_rule(slide, x, y, x + 0.55, y, GOLD, 1.0)
    add_dot(slide, x - 0.04, y - 0.04, 0.08, GOLD, GOLD)
    add_dot(slide, x + 0.51, y - 0.04, 0.08, TEAL, TEAL)


def build(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    # Top identity line: one title, one reading cue, no dense legend.
    add_text(slide, 0.08, 0.06, 0.92, 0.22, "PWSC-CI", size=11, bold=True)
    add_text(slide, 1.10, 0.08, 1.10, 0.18, "motivation", size=9, color=RED, bold=True)
    add_text(slide, 2.30, 0.08, 2.00, 0.18, "two innovations", size=9, color=TEAL, bold=True)
    add_text(slide, 5.64, 0.08, 1.12, 0.18, "risk + answer", size=9, color=MUTED, align=PP_ALIGN.RIGHT)
    add_rule(slide, 0.08, 0.30, W - 0.08, 0.30, SLATE, 1.1)

    # Motivation: two failure cases are carried by visual counterexamples.
    add_box(slide, 0.08, 0.39, 2.05, 2.05, RED_PALE, RGBColor(241, 173, 163))
    add_text(slide, 0.19, 0.45, 1.83, 0.18, "WHY TWO SIGNALS?", size=10, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 0.19, 0.68, 1.83, 0.15, "confidence only", size=9, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_box(slide, 0.20, 0.86, 0.78, 0.28, WHITE, RGBColor(241, 173, 163))
    add_box(slide, 1.23, 0.86, 0.78, 0.28, WHITE, RGBColor(241, 173, 163))
    add_text(slide, 0.24, 0.90, 0.12, 0.12, "A", size=9, color=RED, bold=True)
    add_text(slide, 1.27, 0.90, 0.12, 0.12, "B", size=9, color=RED, bold=True)
    for base in (0.45, 1.48):
        for i in range(4):
            add_box(slide, base + i * 0.105, 0.98, 0.07, 0.07, RED if i == 2 else RGBColor(249, 206, 198), RED, False)
    add_icon_node(slide, 0.37, 1.18, TEAL, fork=False)
    add_icon_node(slide, 1.42, 1.18, RED, fork=True)
    add_text(slide, 0.19, 1.45, 1.83, 0.17, "same q  ≠  same behavior", size=9, color=TEAL, bold=True, align=PP_ALIGN.CENTER)

    add_text(slide, 0.19, 1.66, 1.83, 0.15, "majority only", size=9, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    for i in range(4):
        add_dot(slide, 0.38 + i * 0.16, 1.90, 0.11, TEAL, TEAL)
    add_dot(slide, 1.54, 1.90, 0.14, GOLD, GOLD)
    add_rule(slide, 1.10, 1.96, 1.43, 1.96, GRID, 1.0)
    add_text(slide, 0.19, 2.08, 1.83, 0.15, "count  ≠  target risk", size=9, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 0.18, 2.27, 1.85, 0.16, "quality  ×  structure", size=10, color=SLATE, bold=True, align=PP_ALIGN.CENTER)

    # Innovation 1: quality-weighted task geometry.
    add_box(slide, 2.30, 0.39, 1.67, 2.05, GOLD_PALE, RGBColor(225, 193, 102))
    add_text(slide, 2.39, 0.45, 1.49, 0.16, "INNOVATION 1", size=9, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 2.39, 0.63, 1.49, 0.28, "quality-weighted\ntask geometry", size=9, color=INK, bold=True, align=PP_ALIGN.CENTER)

    # Local quality stream: weakest token -> validity gate -> q_i.
    for i in range(5):
        add_box(slide, 2.46 + i * 0.10, 0.98, 0.07, 0.07, RED if i == 3 else RGBColor(247, 215, 137), RED if i == 3 else GOLD, False)
    add_text(slide, 2.39, 1.10, 0.62, 0.13, "weakest-k", size=9, color=MUTED, align=PP_ALIGN.CENTER)
    add_rule(slide, 2.99, 1.015, 3.13, 1.015, GOLD, 0.9)
    add_dot(slide, 3.12, 0.96, 0.14, TEAL, TEAL)
    add_formula(slide, 3.06, 1.10, 0.26, 0.13, [[("v", None), ("i", "sub")]], size=9, color=TEAL)
    add_formula(slide, 3.00, 1.25, 0.80, 0.18, [[("q", None), ("i", "sub"), ("=c", None), ("i", "sub"), ("v", None), ("i", "sub")]], size=10)

    # Task geometry stream: icons -> positive semidefinite Gram kernel.
    add_domain_icons(slide, 2.45, 1.55)
    add_text(slide, 2.38, 1.78, 0.78, 0.14, "adapters", size=9, color=MUTED, align=PP_ALIGN.CENTER)
    add_matrix(slide, 3.10, 1.54, 2, 2, TEAL, RGBColor(190, 228, 224))
    add_formula(slide, 3.00, 1.88, 0.68, 0.18, [[("K=ΦΦ", None), ("T", "sup")]], size=10)
    add_rule(slide, 3.77, 1.30, 3.77, 1.48, GOLD, 0.9)
    add_rule(slide, 3.77, 1.48, 3.77, 1.68, TEAL, 0.9)
    add_text(slide, 3.67, 2.09, 0.33, 0.14, "mass", size=9, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    add_weight_bar(slide, 3.72, 1.83, GOLD, 0.20, pale=RGBColor(246, 232, 183))
    add_arrow(slide, 3.99, 1.30, 0.12, SLATE)

    # Innovation 2: target-conditioned spectral contrast.
    add_box(slide, 4.14, 0.39, 2.62, 2.05, TEAL_PALE, RGBColor(153, 206, 201))
    add_text(slide, 4.25, 0.45, 2.40, 0.16, "INNOVATION 2", size=9, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 4.25, 0.63, 2.40, 0.16, "target-conditioned spectral contrast", size=9, color=INK, bold=True, align=PP_ALIGN.CENTER)

    # Target / peers are shown as a mass split, not a paragraph.
    add_box(slide, 4.27, 0.87, 0.56, 0.27, GOLD_PALE, RGBColor(225, 193, 102))
    add_text(slide, 4.31, 0.92, 0.48, 0.14, "0  target", size=9, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    for i in range(4):
        add_dot(slide, 4.97 + i * 0.14, 0.95, 0.10, TEAL, TEAL)
    add_text(slide, 4.28, 1.17, 0.52, 0.13, "α", size=9, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 5.08, 1.17, 0.52, 0.13, "β", size=9, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_rule(slide, 5.68, 1.01, 5.84, 1.01, SLATE, 0.9)

    # Two matrix lanes and their spectra make the contrast visible.
    add_text(slide, 5.68, 0.79, 0.34, 0.13, "mix", size=9, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    add_matrix(slide, 5.68, 0.90, 2, 3, GOLD, RGBColor(246, 232, 183))
    add_weight_bar(slide, 6.23, 0.91, GOLD, 0.19, pale=RGBColor(246, 232, 183))
    add_text(slide, 5.68, 1.28, 0.34, 0.13, "peer", size=9, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_matrix(slide, 5.68, 1.39, 2, 3, TEAL, RGBColor(190, 228, 224))
    add_weight_bar(slide, 6.23, 1.40, TEAL, 0.16, pale=RGBColor(190, 228, 224))
    add_text(slide, 6.14, 1.15, 0.48, 0.13, "spectra", size=9, color=MUTED, align=PP_ALIGN.CENTER)

    add_formula(
        slide,
        4.27,
        1.62,
        1.10,
        0.32,
        [
            [("χ", None), ("0", "sub"), ("=S", None), ("mix", "sub")],
            [("−βS", None), ("peer", "sub")],
        ],
        size=10,
    )
    add_formula(
        slide,
        5.54,
        1.62,
        1.10,
        0.32,
        [[("R", None), ("CI", "sub"), ("=clip", None), ("[0,1]", "sub")], [("χ", None), ("0", "sub"), ("/H", None), ("2", "sub"), ("(α)", None)]],
        size=10,
    )

    # The branch is a byproduct, visually separated from the risk score.
    add_box(slide, 4.27, 2.04, 0.91, 0.30, WHITE, RGBColor(225, 193, 102))
    add_formula(slide, 4.31, 2.10, 0.83, 0.18, [[("N=1: 1−q", None), ("0", "sub")]], size=9)
    add_rule(slide, 5.29, 1.82, 5.29, 2.17, TEAL, 0.8)
    add_arrow(slide, 5.29, 2.10, 0.12, TEAL)
    add_box(slide, 5.39, 1.98, 1.20, 0.44, WHITE, RGBColor(153, 206, 201))
    add_text(slide, 5.44, 2.01, 1.10, 0.13, "self-consistent", size=9, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_formula(slide, 5.43, 2.14, 1.12, 0.15, [[("i", None), ("*", "sup"), ("=argmax", None), ("i", "sub")]], size=9)
    add_formula(slide, 5.43, 2.28, 1.12, 0.15, [[("q", None), ("i", "sub"), ("(Kw)", None), ("i", "sub")]], size=9)

    # Production notes remain on slide 2, keeping slide 1 export-ready.
    notes = prs.slides.add_slide(blank)
    notes.background.fill.solid()
    notes.background.fill.fore_color.rgb = WHITE
    add_text(notes, 0.35, 0.35, 6.1, 0.35, "AAAI-27 production notes", size=14, bold=True)
    add_text(
        notes,
        0.40,
        0.95,
        6.0,
        1.1,
        "Final canvas: 17.42 × 6.60 cm\nFigure labels and formulas: 9 pt minimum at final size\nKeep slide 1 at 100% when exporting; use vector PDF or ≥300 dpi raster",
        size=10,
        color=MUTED,
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.output)


if __name__ == "__main__":
    main()
